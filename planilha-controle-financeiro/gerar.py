from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import PieChart, Reference
wb=Workbook()
h=Font(bold=True,color="FFFFFF");fill=PatternFill("solid",fgColor="1F4E78")
cats=["Moradia","Mercado","Transporte","Saúde","Lazer","Educação","Outros"]
ws=wb.active;ws.title="Lançamentos"
ws.append(["Data","Descrição","Tipo","Categoria","Valor (R$)","Mês"])
dados=[("2026-09-01","Salário","Entrada","Outros",3200),("2026-09-02","Aluguel","Saída","Moradia",900),
("2026-09-05","Supermercado","Saída","Mercado",520.4),("2026-09-08","Gasolina","Saída","Transporte",180),
("2026-09-12","Farmácia","Saída","Saúde",76.9),("2026-09-15","Cinema","Saída","Lazer",60),
("2026-09-20","Curso online","Saída","Educação",97),("2026-09-25","Freela","Entrada","Outros",450),
("2026-10-01","Salário","Entrada","Outros",3200),("2026-10-02","Aluguel","Saída","Moradia",900),
("2026-10-06","Supermercado","Saída","Mercado",610.15),("2026-10-09","Ônibus","Saída","Transporte",95)]
import datetime
for i,(d,de,t,c,v) in enumerate(dados,start=2):
    ws.append([datetime.date.fromisoformat(d),de,t,c,v,f'=IF(A{i}="","",DATE(YEAR(A{i}),MONTH(A{i}),1))'])
for i in range(2,201):
    ws[f"A{i}"].number_format="DD/MM/YYYY";ws[f"F{i}"].number_format="MM/YYYY";ws[f"E{i}"].number_format='"R$" #,##0.00'
    if i>len(dados)+1: ws[f"F{i}"]=f'=IF(A{i}="","",DATE(YEAR(A{i}),MONTH(A{i}),1))'
for c in ws[1]: c.font=h;c.fill=fill
for col,w in zip("ABCDEF",[12,28,10,14,14,10]): ws.column_dimensions[col].width=w
dv=DataValidation(type="list",formula1='"Entrada,Saída"');ws.add_data_validation(dv);dv.add("C2:C200")
dv2=DataValidation(type="list",formula1='"'+",".join(cats)+'"');ws.add_data_validation(dv2);dv2.add("D2:D200")
ws.freeze_panes="A2"
r=wb.create_sheet("Resumo")
r["A1"]="Resumo financeiro";r["A1"].font=Font(bold=True,size=14)
r["A2"]="Mês de referência";r["B2"]=datetime.date(2026,9,1);r["B2"].number_format="MM/YYYY";r["C2"]="(troque o mês aqui)"
r["A3"]="Total de entradas";r["B3"]='=SUMIFS(Lançamentos!E:E,Lançamentos!C:C,"Entrada",Lançamentos!F:F,B2)'
r["A4"]="Total de saídas";r["B4"]='=SUMIFS(Lançamentos!E:E,Lançamentos!C:C,"Saída",Lançamentos!F:F,B2)'
r["A5"]="Saldo";r["B5"]="=B3-B4";r["A5"].font=Font(bold=True)
r["A7"]="Categoria";r["B7"]="Gasto";r["C7"]="Orçamento";r["D7"]="% usado"
for c in r[7]: c.font=h;c.fill=fill
orc=[1000,600,250,150,150,120,200]
for i,(c,o) in enumerate(zip(cats,orc),start=8):
    r[f"A{i}"]=c;r[f"B{i}"]=f'=SUMIFS(Lançamentos!E:E,Lançamentos!C:C,"Saída",Lançamentos!D:D,A{i},Lançamentos!F:F,$B$2)'
    r[f"C{i}"]=o;r[f"D{i}"]=f'=IF(C{i}=0,0,B{i}/C{i})';r[f"D{i}"].number_format="0%"
for c in ["B3","B4","B5"]+[f"{x}{i}" for x in "BC" for i in range(8,15)]: r[c].number_format='"R$" #,##0.00'
r.conditional_formatting.add("D8:D14",CellIsRule(operator="greaterThan",formula=["1"],fill=PatternFill("solid",fgColor="F8CBAD")))
r.conditional_formatting.add("B5",CellIsRule(operator="lessThan",formula=["0"],font=Font(color="C00000",bold=True)))
r.column_dimensions["A"].width=20
for x in "BCD": r.column_dimensions[x].width=14
ch=PieChart();ch.title="Gastos por categoria"
ch.add_data(Reference(r,min_col=2,min_row=7,max_row=14),titles_from_data=True)
ch.set_categories(Reference(r,min_col=1,min_row=8,max_row=14));r.add_chart(ch,"F3")
wb.save("controle-financeiro.xlsx")
